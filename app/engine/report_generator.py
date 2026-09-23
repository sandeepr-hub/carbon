import re
import os
import datetime
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY

from sqlalchemy.orm import Session
from app.models.assessment import Assessment, AssessmentSummary
from app.models.campus import Campus, Building
from app.models.carbon import CarbonCalculation, CarbonReductionContribution, EmissionFactor
from app.engine.aggregator import get_building_wise_results, get_source_breakdown
from app.engine.recommendation_engine import generate_campus_recommendations
from app.config import EXPORT_DIR

def generate_pdf_report(db: Session, assessment_id: int) -> str:
    assessment = db.query(Assessment).filter(Assessment.id == assessment_id).first()
    if not assessment:
        raise ValueError("Assessment not found")

    campus = assessment.campus
    summary = assessment.summary
    buildings_data = get_building_wise_results(db, assessment_id)
    sources_data = get_source_breakdown(db, assessment_id)
    recommendations = generate_campus_recommendations(db, assessment_id)
    reductions = db.query(CarbonReductionContribution).filter(
        CarbonReductionContribution.assessment_id == assessment_id
    ).all()

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^a-zA-Z0-9_-]", "_", campus.name)
    file_name = f"Campus_Carbon_Report_{safe_name}_{assessment.reporting_year}_{timestamp}.pdf"
    file_path = EXPORT_DIR / file_name

    doc = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'CoverTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#064e3b'), # Dark Emerald
        alignment=TA_CENTER
    )
    subtitle_style = ParagraphStyle(
        'CoverSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#047857'),
        alignment=TA_CENTER
    )
    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#0f766e'),
        spaceBefore=14,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#1e293b')
    )
    callout_style = ParagraphStyle(
        'Callout',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155')
    )

    story = []

    # --- Header Banner ---
    story.append(Paragraph("CAMPUS CARBON", title_style))
    story.append(Paragraph("Universal Campus Carbon Footprint Assessment & Management Platform (v2.0)", subtitle_style))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#059669'), spaceBefore=4, spaceAfter=12))

    # --- Campus & Assessment Profile ---
    profile_data = [
        [
            Paragraph("<b>Organization:</b>", body_style), Paragraph(campus.organization.name if campus.organization else "N/A", body_style),
            Paragraph("<b>Reporting Year:</b>", body_style), Paragraph(str(assessment.reporting_year), body_style)
        ],
        [
            Paragraph("<b>Campus Name:</b>", body_style), Paragraph(campus.name, body_style),
            Paragraph("<b>Assessment Status:</b>", body_style), Paragraph(assessment.status, body_style)
        ],
        [
            Paragraph("<b>Campus Type:</b>", body_style), Paragraph(campus.campus_type, body_style),
            Paragraph("<b>Total Population:</b>", body_style), Paragraph(f"{campus.population:,} occupants", body_style)
        ],
        [
            Paragraph("<b>Location:</b>", body_style), Paragraph(f"{campus.city}, {campus.state}, {campus.country}", body_style),
            Paragraph("<b>Built-up Area:</b>", body_style), Paragraph(f"{campus.total_built_up_area_sqm:,.1f} m² ({campus.area} {campus.area_unit})", body_style)
        ],
        [
            Paragraph("<b>Standard / Protocol:</b>", body_style), Paragraph(assessment.methodology, body_style),
            Paragraph("<b>Generated On:</b>", body_style), Paragraph(datetime.datetime.now().strftime("%d %B %Y, %H:%M"), body_style)
        ]
    ]

    t_profile = Table(profile_data, colWidths=[100, 160, 100, 160])
    t_profile.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('PADDING', (0,0), (-1,-1), 4),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(t_profile)
    story.append(Spacer(1, 14))

    # --- Pillar Hero Scorecard ---
    gross_val = summary.gross_emissions if summary else 0.0
    red_val = summary.eligible_reductions if summary else 0.0
    net_val = summary.net_footprint if summary else 0.0

    scorecard_data = [
        [
            Paragraph("<font size='10'><b>1. GROSS / POSITIVE EMISSIONS</b></font>", ParagraphStyle('SC1', alignment=TA_CENTER, textColor=colors.HexColor('#991b1b'))),
            Paragraph("<font size='10'><b>2. ELIGIBLE REDUCTIONS & REMOVALS</b></font>", ParagraphStyle('SC2', alignment=TA_CENTER, textColor=colors.HexColor('#065f46'))),
            Paragraph("<font size='10'><b>3. NET CARBON FOOTPRINT</b></font>", ParagraphStyle('SC3', alignment=TA_CENTER, textColor=colors.HexColor('#1e40af')))
        ],
        [
            Paragraph(f"<font size='18'><b>{gross_val:,.2f}</b></font><br/><font size='9'>tCO2e</font>", ParagraphStyle('SCV1', alignment=TA_CENTER, textColor=colors.HexColor('#b91c1c'))),
            Paragraph(f"<font size='18'><b>-{red_val:,.2f}</b></font><br/><font size='9'>tCO2e</font>", ParagraphStyle('SCV2', alignment=TA_CENTER, textColor=colors.HexColor('#047857'))),
            Paragraph(f"<font size='18'><b>{net_val:,.2f}</b></font><br/><font size='9'>tCO2e</font>", ParagraphStyle('SCV3', alignment=TA_CENTER, textColor=colors.HexColor('#1d4ed8')))
        ]
    ]

    t_scorecard = Table(scorecard_data, colWidths=[170, 180, 170])
    t_scorecard.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#fef2f2')), # light red
        ('BACKGROUND', (1,0), (1,-1), colors.HexColor('#ecfdf5')), # light green
        ('BACKGROUND', (2,0), (2,-1), colors.HexColor('#eff6ff')), # light blue
        ('BOX', (0,0), (0,-1), 1.5, colors.HexColor('#fca5a5')),
        ('BOX', (1,0), (1,-1), 1.5, colors.HexColor('#6ee7b7')),
        ('BOX', (2,0), (2,-1), 1.5, colors.HexColor('#93c5fd')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_scorecard)
    story.append(Spacer(1, 14))

    # --- GHG Scope Classification & Intensity Indicators ---
    story.append(Paragraph("1. Executive Carbon Accounting Breakdown", h1_style))
    
    scope_data = [
        ["Scope Classification", "Activity Description", "Emissions (tCO2e)", "Share of Gross (%)"],
        ["Scope 1 (Direct)", "Stationary DG diesel, canteen LPG, campus fleet combustion, refrigerants", f"{summary.scope1:,.2f}" if summary else "0.00", f"{(summary.scope1/gross_val*100):.1f}%" if summary and gross_val > 0 else "0%"],
        ["Scope 2 (Indirect - Electricity)", "Purchased grid electricity & EV charging consumption", f"{summary.scope2:,.2f}" if summary else "0.00", f"{(summary.scope2/gross_val*100):.1f}%" if summary and gross_val > 0 else "0%"],
        ["Scope 3 (Other Indirect)", "Commuting, solid waste disposal, freshwater supply, wastewater STP", f"{summary.scope3:,.2f}" if summary else "0.00", f"{(summary.scope3/gross_val*100):.1f}%" if summary and gross_val > 0 else "0%"],
        ["Total Gross Emissions", "All combined operational activity emissions", f"{gross_val:,.2f}", "100.0%"],
        ["Eligible Removals / Reductions", "Tree carbon sequestration, verified EV savings, waste diversion", f"-{red_val:,.2f}", f"-{(red_val/gross_val*100):.1f}%" if gross_val > 0 else "0%"],
        ["Official Net Carbon Footprint", "Gross Emissions − Eligible Reductions & Removals", f"{net_val:,.2f}", f"{(net_val/gross_val*100):.1f}%" if gross_val > 0 else "0%"]
    ]

    t_scopes = Table(scope_data, colWidths=[130, 230, 90, 70])
    t_scopes.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f766e')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8.5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0,1), (-1,-4), [colors.white, colors.HexColor('#f8fafc')]),
        ('BACKGROUND', (0,-3), (-1,-3), colors.HexColor('#fee2e2')), # Gross row
        ('FONTNAME', (0,-3), (-1,-3), 'Helvetica-Bold'),
        ('BACKGROUND', (0,-2), (-1,-2), colors.HexColor('#d1fae5')), # Reductions row
        ('FONTNAME', (0,-2), (-1,-2), 'Helvetica-Bold'),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#dbeafe')), # Net row
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('ALIGN', (2,0), (3,-1), 'RIGHT'),
        ('PADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_scopes)
    story.append(Spacer(1, 10))

    # Intensity Metrics Table
    intensity_data = [
        [
            Paragraph(f"<b>Gross Intensity / Person:</b><br/>{summary.gross_per_person:.4f} tCO2e/person", body_style),
            Paragraph(f"<b>Net Intensity / Person:</b><br/>{summary.net_per_person:.4f} tCO2e/person", body_style),
            Paragraph(f"<b>Gross Intensity / Area:</b><br/>{summary.gross_per_area*1000:.2f} kgCO2e/m²", body_style),
            Paragraph(f"<b>Renewable Energy Share:</b><br/>{summary.renewable_percentage:.1f}%", body_style),
            Paragraph(f"<b>Waste Diversion Rate:</b><br/>{summary.waste_diversion_rate:.1f}%", body_style)
        ]
    ]
    t_intensity = Table(intensity_data, colWidths=[104, 104, 104, 104, 104])
    t_intensity.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f1f5f9')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#94a3b8')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_intensity)
    story.append(Spacer(1, 14))

    # --- Building-wise Performance Table ---
    story.append(Paragraph("2. Building-wise Carbon Performance Leaderboard", h1_style))
    bld_headers = ["Code", "Building Name", "Type", "Area (m²)", "Occ.", "Gross (t)", "Red. (t)", "Net (t)", "tCO2e/p", "kg/m²"]
    bld_rows = [bld_headers]

    for b in buildings_data:
        bld_rows.append([
            b["building_code"],
            b["name"][:25] + "..." if len(b["name"]) > 25 else b["name"],
            b["building_type"],
            f"{b['built_up_area_sqm']:,.0f}",
            str(b["occupancy"]),
            f"{b['gross_tco2e']:,.2f}",
            f"{b['reductions_tco2e']:,.2f}",
            f"{b['net_tco2e']:,.2f}",
            f"{b['co2e_per_person']:.2f}",
            f"{b['co2e_per_sqm']*1000:.1f}"
        ])

    t_bld = Table(bld_rows, colWidths=[45, 120, 55, 45, 30, 45, 45, 45, 45, 45])
    t_bld.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 7.5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
        ('ALIGN', (3,0), (-1,-1), 'RIGHT'),
        ('PADDING', (0,0), (-1,-1), 3.5),
    ]))
    story.append(t_bld)
    story.append(Spacer(1, 14))

    # --- Negative & Removal Contributions Ledger ---
    story.append(Paragraph("3. Negative / Eligible Reduction Contributions Ledger", h1_style))
    red_headers = ["Category", "Activity & Intervention", "Methodology & Baseline", "Status", "Avoided / Sequestered"]
    red_rows = [red_headers]

    for r in reductions:
        red_rows.append([
            r.category,
            r.activity[:35] + "..." if len(r.activity) > 35 else r.activity,
            r.methodology[:40] + "..." if len(r.methodology) > 40 else r.methodology,
            r.eligibility_status,
            f"-{r.reduction_co2e:,.2f} tCO2e"
        ])

    if len(reductions) == 0:
        red_rows.append(["None", "No reduction/removal activities logged yet", "-", "-", "0.00 tCO2e"])

    t_red = Table(red_rows, colWidths=[80, 160, 160, 55, 65])
    t_red.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#065f46')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 7.5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f0fdf4')]),
        ('ALIGN', (-1,0), (-1,-1), 'RIGHT'),
        ('PADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_red)
    story.append(Spacer(1, 14))

    # --- Prioritized Action Plan & Recommendations ---
    story.append(Paragraph("4. Prioritized Recommendations & Decarbonization Action Plan", h1_style))
    rec_headers = ["Priority", "Category", "Intervention Action Title", "Intervention Type", "Est. Impact"]
    rec_rows = [rec_headers]

    for rec in recommendations[:5]:
        rec_rows.append([
            rec["priority"],
            rec["category"],
            rec["title"][:40] + "..." if len(rec["title"]) > 40 else rec["title"],
            rec["intervention_type"],
            f"-{rec['estimated_reduction_pct']:.0f}%"
        ])

    t_rec = Table(rec_rows, colWidths=[55, 75, 230, 110, 50])
    t_rec.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f766e')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
        ('ALIGN', (-1,0), (-1,-1), 'CENTER'),
        ('PADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_rec)
    story.append(Spacer(1, 14))

    # --- Methodological & Anti-Double-Counting Safeguards Disclosure ---
    story.append(Paragraph("5. Accounting Boundaries & Anti-Double-Counting Disclosures", h1_style))
    disclosure_text = (
        "<b>Methodological Boundary:</b> This assessment is executed under the operational control consolidation approach in strict accordance with the GHG Protocol Corporate Accounting and Reporting Standard and ISO 14064-1:2018 specifications.<br/>"
        "<b>Anti-Double-Counting Safeguards:</b><br/>"
        "1. <i>On-site Solar Generation:</i> Solar energy consumed directly on campus displaces grid power and is reflected exclusively in reduced Scope 2 emissions. It is NOT subtracted a second time as a separate negative carbon offset.<br/>"
        "2. <i>EV Fleet Accounting:</i> Operational EV charging electricity is quantified under Scope 2 electricity. Avoided transport reductions are calculated separately against an explicit ICE baseline.<br/>"
        "3. <i>Carbon Sequestration:</i> Tree and vegetation biomass sequestration is modeled based on Forest Survey of India & IPCC GPG factors (22.0 kg CO2/tree/year) and reported in the eligible removals ledger."
    )
    story.append(Paragraph(disclosure_text, callout_style))
    story.append(Spacer(1, 14))

    doc.build(story)
    return str(file_path)

def generate_excel_report(db: Session, assessment_id: int) -> str:
    assessment = db.query(Assessment).filter(Assessment.id == assessment_id).first()
    if not assessment:
        raise ValueError("Assessment not found")

    campus = assessment.campus
    summary = assessment.summary
    buildings_data = get_building_wise_results(db, assessment_id)
    calcs = db.query(CarbonCalculation).filter(CarbonCalculation.assessment_id == assessment_id).all()
    reductions = db.query(CarbonReductionContribution).filter(CarbonReductionContribution.assessment_id == assessment_id).all()
    factors = db.query(EmissionFactor).all()

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^a-zA-Z0-9_-]", "_", campus.name)
    file_name = f"Campus_Carbon_Report_{safe_name}_{assessment.reporting_year}_{timestamp}.xlsx"
    file_path = EXPORT_DIR / file_name

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"

    # Styling helper
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    title_font = Font(name="Calibri", size=14, bold=True, color="064E3B")

    # Tab 1: Executive Summary
    ws_summary["A1"] = "CAMPUS CARBON ASSESSMENT REPORT"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"{campus.name} — Reporting Year {assessment.reporting_year}"
    ws_summary["A2"].font = bold_font

    summary_rows = [
        ["Metric", "Value", "Unit", "Notes"],
        ["Gross / Positive Emissions", summary.gross_emissions if summary else 0, "tCO2e", "Total Scope 1 + Scope 2 + Scope 3 operational emissions"],
        ["Eligible Reductions / Removals", summary.eligible_reductions if summary else 0, "tCO2e", "Verified carbon sequestration & avoided emissions"],
        ["Net Carbon Footprint", summary.net_footprint if summary else 0, "tCO2e", "Gross Emissions − Eligible Reductions"],
        ["Scope 1 (Direct)", summary.scope1 if summary else 0, "tCO2e", "Stationary fuels, DG, mobile fleet, refrigerants"],
        ["Scope 2 (Electricity)", summary.scope2 if summary else 0, "tCO2e", "Purchased grid electricity & EV charging"],
        ["Scope 3 (Indirect)", summary.scope3 if summary else 0, "tCO2e", "Commuting, solid waste, water supply & STP"],
        ["Total Campus Population", summary.population if summary else 0, "Occupants", "Students, faculty, staff, residents"],
        ["Gross CO2e per Person", summary.gross_per_person if summary else 0, "tCO2e/person", "Gross emissions divided by population"],
        ["Net CO2e per Person", summary.net_per_person if summary else 0, "tCO2e/person", "Net footprint divided by population"],
        ["Total Built-up Area", summary.built_up_area_sqm if summary else 0, "Sq. Meters", "Standardized built-up area across all buildings"],
        ["Gross CO2e per m²", summary.gross_per_area if summary else 0, "tCO2e/m²", "Gross emissions per square meter built-up area"],
        ["Renewable Energy Share", summary.renewable_percentage if summary else 0, "%", "Solar/clean energy generation as % of total power"],
        ["Waste Diversion Rate", summary.waste_diversion_rate if summary else 0, "%", "Waste composted/recycled as % of total waste"]
    ]

    for r_idx, row in enumerate(summary_rows, start=4):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            if r_idx == 4:
                cell.fill = header_fill
                cell.font = header_font
            elif c_idx == 1:
                cell.font = bold_font

    # Tab 2: Building Performance
    ws_bld = wb.create_sheet(title="Building Performance")
    bld_headers = ["Building Code", "Building Name", "Type", "Floors", "Built-up Area (m²)", "Occupancy", "Gross (tCO2e)", "Reductions (tCO2e)", "Net (tCO2e)", "tCO2e/Person", "tCO2e/m²", "Scope 1 (tCO2e)", "Scope 2 (tCO2e)", "Scope 3 (tCO2e)"]
    ws_bld.append(bld_headers)
    for c_idx in range(1, len(bld_headers) + 1):
        ws_bld.cell(row=1, column=c_idx).fill = header_fill
        ws_bld.cell(row=1, column=c_idx).font = header_font

    for b in buildings_data:
        ws_bld.append([
            b["building_code"], b["name"], b["building_type"], b["floors"],
            b["built_up_area_sqm"], b["occupancy"], b["gross_tco2e"],
            b["reductions_tco2e"], b["net_tco2e"], b["co2e_per_person"],
            b["co2e_per_sqm"], b["scope1_tco2e"], b["scope2_tco2e"], b["scope3_tco2e"]
        ])

    # Tab 3: Detailed Calculations Log (Gross)
    ws_calc = wb.create_sheet(title="Gross Calculations Log")
    calc_headers = ["ID", "Category", "Activity Description", "Scope", "Quantity", "Unit", "Emission Factor", "Factor Unit", "Factor Source", "Emissions (tCO2e)", "Calculation Method", "Data Quality"]
    ws_calc.append(calc_headers)
    for c_idx in range(1, len(calc_headers) + 1):
        ws_calc.cell(row=1, column=c_idx).fill = header_fill
        ws_calc.cell(row=1, column=c_idx).font = header_font

    for c in calcs:
        ws_calc.append([
            c.id, c.category, c.activity, c.scope, c.quantity, c.unit,
            c.emission_factor, c.emission_factor_unit, c.emission_factor_source,
            c.emissions_co2e, c.calculation_method, c.data_quality
        ])

    # Tab 4: Reductions & Removals Ledger
    ws_red = wb.create_sheet(title="Reductions & Removals")
    red_headers = ["ID", "Category", "Activity", "Contribution Type", "Baseline Quantity", "Baseline Unit", "Project Quantity", "Project Unit", "Emission Factor", "Gross Baseline (tCO2e)", "Project Emissions (tCO2e)", "Reduction (tCO2e)", "Methodology", "Status", "Assumptions"]
    ws_red.append(red_headers)
    for c_idx in range(1, len(red_headers) + 1):
        ws_red.cell(row=1, column=c_idx).fill = header_fill
        ws_red.cell(row=1, column=c_idx).font = header_font

    for r in reductions:
        ws_red.append([
            r.id, r.category, r.activity, r.contribution_type,
            r.baseline_quantity, r.baseline_unit, r.project_quantity, r.project_unit,
            r.emission_factor, r.gross_baseline_emissions, r.project_emissions,
            r.reduction_co2e, r.methodology, r.eligibility_status, r.assumptions
        ])

    # Tab 5: Emission Factors Library
    ws_ef = wb.create_sheet(title="Emission Factors Library")
    ef_headers = ["ID", "Category", "Activity", "Fuel Type", "Unit", "Factor", "Factor Unit", "Scope", "Source", "Ref Year", "Region", "Notes"]
    ws_ef.append(ef_headers)
    for c_idx in range(1, len(ef_headers) + 1):
        ws_ef.cell(row=1, column=c_idx).fill = header_fill
        ws_ef.cell(row=1, column=c_idx).font = header_font

    for ef in factors:
        ws_ef.append([
            ef.id, ef.category, ef.activity, ef.fuel_type, ef.unit,
            ef.factor, ef.factor_unit, ef.scope, ef.source, ef.reference_year,
            ef.region, ef.notes
        ])

    wb.save(str(file_path))
    return str(file_path)

def generate_csv_report(db: Session, assessment_id: int) -> str:
    buildings_data = get_building_wise_results(db, assessment_id)
    assessment = db.query(Assessment).filter(Assessment.id == assessment_id).first()
    campus = assessment.campus

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^a-zA-Z0-9_-]", "_", campus.name)
    file_name = f"Campus_Carbon_Buildings_{safe_name}_{assessment.reporting_year}_{timestamp}.csv"
    file_path = EXPORT_DIR / file_name

    df = pd.DataFrame(buildings_data)
    df.to_csv(str(file_path), index=False)
    return str(file_path)
